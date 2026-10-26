"""Build the detailed ten-slide Phase 1 PowerPoint and a matching PDF.

Install reports/requirements.txt. Metrics are read from the completed experiment;
no training, parameter selection, or experiment output is changed by this script.
"""

from __future__ import annotations

import csv
import json
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab import rl_config
from PIL import Image

from build_presentation_figures import build_figures


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports/presentation"
RUN = ROOT / "artifacts/phase1"
TEAM = [
    ("Samarth Khera", "2401010407"),
    ("Pratham Malhotra", "2401010343"),
    ("Adityaraj Pal", "2401010042"),
    ("Tattva Rajput", "2401010478"),
    ("Himanshu", "2401010188"),
]
NAVY = "102633"
INK = "152E3A"
MUTED = "526774"
PAPER = "F4F5F0"
WHITE = "FFFFFF"
LINE = "D4DDD8"
MINT = "3DE0B1"
GREEN = "087E64"
ORANGE = "DC813F"
SOFT = "E5EEE7"
BLUE = "3779AA"
WIDTH, HEIGHT = 1280, 720
TOTAL_SLIDES = 10


def records(name):
    with (RUN / name).open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def register_fonts():
    pairs = [
        (os.environ.get("AIRCAST_FONT_REGULAR"), os.environ.get("AIRCAST_FONT_BOLD")),
        ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ]
    for regular, bold in pairs:
        if regular and bold and Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("Deck", regular))
            pdfmetrics.registerFont(TTFont("DeckBold", bold))
            return "Arial" if "Arial" in regular or "arial" in regular else "DejaVu Sans"
    raise RuntimeError("Set AIRCAST_FONT_REGULAR and AIRCAST_FONT_BOLD to TrueType font paths")


class Deck:
    """Draw the same measured text lines and vector shapes into both formats."""

    def __init__(self):
        OUTPUT.mkdir(parents=True, exist_ok=True)
        # Store compressed images as binary streams, avoiding ASCII85 expansion.
        rl_config.useA85 = False
        self.font = register_fonts()
        self.prs = Presentation()
        self.prs.slide_width = Inches(WIDTH / 72)
        self.prs.slide_height = Inches(HEIGHT / 72)
        self.prs.core_properties.title = "Aircast | Phase 1: Air Quality Prediction"
        self.prs.core_properties.subject = "One-hour PM2.5 forecasting: completed classical ML work"
        self.prs.core_properties.author = "; ".join(name for name, _ in TEAM)
        self.prs.core_properties.keywords = "PM2.5, Phase 1, XGBoost, Ridge, chronological backtesting"
        self.pdf = canvas.Canvas(str(OUTPUT / "slides.pdf"), pagesize=(WIDTH, HEIGHT))
        self.pdf.setTitle(self.prs.core_properties.title)
        self.pdf.setAuthor(self.prs.core_properties.author)
        self.slide = None
        self.page = 0
        self.notes = []

    def rect(self, x, y, w, h, fill, radius=0):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        if radius:
            shape.adjustments[0] = min(radius / min(w, h), 0.25)
        self.pdf.setFillColor(HexColor("#" + fill))
        if radius:
            self.pdf.roundRect(x, HEIGHT-y-h, w, h, radius, fill=1, stroke=0)
        else:
            self.pdf.rect(x, HEIGHT-y-h, w, h, fill=1, stroke=0)

    def line(self, x1, y1, x2, y2, color=LINE, width=1):
        from pptx.enum.shapes import MSO_CONNECTOR
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        shape.line.color.rgb = RGBColor.from_string(color)
        shape.line.width = Pt(width)
        self.pdf.setStrokeColor(HexColor("#" + color))
        self.pdf.setLineWidth(width)
        self.pdf.line(x1, HEIGHT-y1, x2, HEIGHT-y2)

    def figure(self, name, x, y, w, h):
        path = OUTPUT / "figures" / f"{name}.png"
        with Image.open(path) as graphic:
            iw, ih = graphic.size
        scale = min(w / iw, h / ih)
        fw, fh = iw * scale, ih * scale
        left, top = x + (w-fw)/2, y + (h-fh)/2
        self.slide.shapes.add_picture(str(path), Pt(left), Pt(top), width=Pt(fw), height=Pt(fh))
        self.pdf.drawImage(ImageReader(str(path)), left, HEIGHT-top-fh, fw, fh, mask="auto")

    def panel(self, x, y, w, h, title, subtitle=None, fill=WHITE):
        self.rect(x, y, w, h, fill, 10)
        self.text(title, x+20, y+17, w-40, 22, INK, True)
        if subtitle:
            self.text(subtitle, x+20, y+49, w-40, 14, MUTED)

    def stat(self, x, y, w, number, label, detail=None):
        self.rect(x, y, w, 91, WHITE, 9)
        self.text(number, x+20, y+10, w-40, 32, GREEN, True)
        self.text(label, x+20, y+52, w-40, 16, INK, True)
        if detail:
            self.text(detail, x+20, y+73, w-40, 12, MUTED)

    def text(self, content, x, y, w, size=18, color=INK, bold=False, leading=1.22, align="left"):
        font = "DeckBold" if bold else "Deck"
        lines = []
        for paragraph in str(content).split("\n"):
            line = ""
            for word in paragraph.split():
                candidate = (line + " " + word).strip()
                if pdfmetrics.stringWidth(candidate, font, size) > w and line:
                    lines.append(line)
                    line = word
                else:
                    line = candidate
            lines.append(line)
        line_h = size * leading
        for i, value in enumerate(lines):
            top = y + i * line_h
            if top + size > HEIGHT - 8:
                raise ValueError(f"Slide {self.page}: text outside content area: {value}")
            shape = self.slide.shapes.add_textbox(Pt(x), Pt(top), Pt(w), Pt(line_h + 3))
            tf = shape.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            p = tf.paragraphs[0]
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = 1.0
            run = p.add_run()
            run.text = value
            run.font.name = self.font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
            self.pdf.setFont(font, size)
            self.pdf.setFillColor(HexColor("#" + color))
            baseline = HEIGHT - top - size * 0.91
            if align == "center":
                self.pdf.drawCentredString(x+w/2, baseline, value)
            elif align == "right":
                self.pdf.drawRightString(x+w, baseline, value)
            else:
                self.pdf.drawString(x, baseline, value)
        return len(lines) * line_h

    def start(self, section, title=None, subtitle=None, dark=False, source=""):
        if self.page:
            self.pdf.showPage()
        self.page += 1
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(0, 0, WIDTH, HEIGHT, NAVY if dark else PAPER)
        self.text("AIRCAST", 56, 34, 170, 17, MINT if dark else GREEN, True)
        self.text(section.upper(), 264, 37, 900, 13, "B8CECE" if dark else MUTED, True)
        self.rect(56, 72, 1168, 1, "2C4753" if dark else LINE)
        if title:
            self.text(title, 56, 102, 1168, 40, WHITE if dark else INK, True)
        if subtitle:
            self.text(subtitle, 56, 159, 1168, 18, "B8CECE" if dark else MUTED)
        self.line(56, 670, 1224, 670, "2C4753" if dark else LINE)
        self.text(source or "CSAI PROJECT 19  /  PHASE 1 · MACHINE LEARNING", 56, 686, 1070, 11, "B8CECE" if dark else MUTED)
        self.text(f"{self.page:02d} / {TOTAL_SLIDES:02d}", 1144, 685, 80, 13, "B8CECE" if dark else MUTED, align="right")

    def note(self, title, content):
        self.slide.notes_slide.notes_text_frame.text = content
        self.notes.append(f"## {self.page}. {title}\n\n{content}\n")

    def save(self):
        if self.page != TOTAL_SLIDES:
            raise ValueError(f"The Phase 1 deck must contain exactly {TOTAL_SLIDES} slides")
        self.prs.save(OUTPUT / "Aircast_Phase1_Presentation.pptx")
        self.pdf.save()
        (OUTPUT / "presenter_guide.md").write_text(
            "# Phase 1 presentation: ten-slide detailed speaker guide\n\n"
            "The PowerPoint contains the same notes. Aim for 12–15 minutes; all errors and concentrations are in µg/m³. "
            "Do not describe R² as percentage accuracy or the historical demo as live air quality. "
            "No individual contribution split has been invented.\n\n" + "\n".join(self.notes), encoding="utf-8")



def main():
    chart_audit = build_figures()
    manifest = json.loads((RUN / "run_manifest.json").read_text())
    metrics = {(r["phase"], r["model"]): r for r in records("metrics.csv")}
    scores = {name: metrics[("holdout", name)] for name in ("persistence", "ridge", "xgboost", "seasonal_naive")}
    selected = manifest["selected_parameters"]["xgboost"]
    n = int(scores["xgboost"]["n"])
    improvement = {name: 100*(float(scores["persistence"]["rmse"])-float(scores[name]["rmse"]))/float(scores["persistence"]["rmse"]) for name in ("ridge", "xgboost")}
    d = Deck()

    # 1. A complete research question, with the Phase 1 work visible on the slide.
    d.start("CSAI Project 19 / Phase 1 research presentation", dark=True)
    d.text("Air quality,\none hour ahead.", 56, 119, 745, 68, WHITE, True, 1.08)
    d.text("Can pollution history and weather improve on simply repeating the latest reading?", 60, 294, 660, 27, "BED1D2")
    d.rect(825, 124, 399, 351, "1B3947", 14)
    d.text("THE PHASE 1 WORKFLOW", 850, 148, 350, 16, MINT, True)
    stages = ["01  Literature + problem framing", "02  Data audit + training-only EDA", "03  Features + four forecasting models", "04  Selection + rolling backtests", "05  Diagnostics + Aircast demonstration"]
    for i, item in enumerate(stages):
        d.text(item, 850, 199+i*49, 350, 18, WHITE)
        if i < 4: d.line(850, 233+i*49, 1194, 233+i*49, "2C4A57", 0.7)
    for x, value, label in [(60, "43,824", "hourly records"), (310, "4", "models compared"), (560, "+1 hour", "forecast horizon")]:
        d.text(value, x, 461, 220, 42, MINT, True)
        d.text(label, x, 517, 220, 18, "BED1D2")
    d.text("UCI Beijing · 2010–2014 · PM2.5 concentration in µg/m³", 60, 598, 1140, 20, "BED1D2")
    d.text("Measured outcome: modest pooled gains, a near tie with Ridge, and unresolved peak errors.", 60, 632, 1140, 17, MINT)
    d.note("Research question and complete Phase 1 scope", "We estimate the observed PM2.5 concentration one hour ahead. After receiving the 10:00 reading, predict 11:00 using pollution and measured weather available through 10:00. Future calendar values are known, but measured target-hour weather is not. The output is concentration in µg/m³. The question is the incremental benefit beyond persistence, not whether a complex model can achieve a large R². Phase 1 delivered literature review, a validated hourly archive, training-only EDA, causal feature engineering, persistence and daily seasonal naive, Ridge and XGBoost, chronological selection, rolling holdout backtests, diagnostics, saved inference, notebooks/CLI, Aircast, tests and documentation. The figures in this deck are built from the completed real-data run, not synthetic fixtures. Separate deep-learning or hybrid modules do not add results to the Phase 1 tables. We will show both gains and limitations, including the strong validation baselines and the difficult high-pollution subset.")

    # 2. Audit and cleaning: two large, fully labelled graphs.
    d.start("01 / Dataset and cleaning", "Preserve the clock, the gaps, and the peaks", "UCI Beijing PM2.5 · January 2010–December 2014 · pollution and airport weather use different sites.", source="SOURCE: UCI 381; supplied CSV audit; concentration distribution uses 2010–2012 only.")
    for x, value, label, detail in [(56,"43,824","Regular hourly timestamps","0 duplicates; 0 absent timestamp rows"), (454,"2,067","Missing concentrations","4.72% of the archive; weather complete"), (852,"155 h","Longest missing run","214 concentration-gap runs in total")]:
        d.stat(x, 199, 372, value, label, detail)
    d.panel(56, 309, 572, 302, "Training concentration distribution")
    d.figure("training_distribution", 72, 355, 540, 242)
    d.panel(652, 309, 572, 302, "Missingness varies by year")
    d.figure("missingness_by_year", 668, 355, 540, 242)
    d.text("Cleaning: validate schema and dates → preserve hourly slots → retain valid peaks → exclude row ID", 60, 631, 1155, 18, GREEN, True)
    d.note("Dataset, gaps and cleaning decisions", "The original UCI 381 file has 43,824 hourly rows, sorted unique timestamps, no absent timestamp rows and 2,067 missing PM2.5 measurements (4.72%). Numeric weather is complete. The gap audit finds 214 concentration runs, with a maximum length of 155 hours. We leave missing measurements in their hourly slots: deleting the 10:00 row would incorrectly make a one-row shift at 11:00 point to 09:00. The loader validates dates, numeric fields and duplicate timestamps, sorts and regularizes before features are built, and removes the No identifier. No target interpolation is used for scoring. Valid high concentrations remain: a statistical outlier is not automatically a recording error. The histogram uses only 2010–2012, with 24,418 measured values, mean 97.82, median 73, range 0–994, skewness 1.70 and p90 218 µg/m³. The missingness plot is a descriptive full-archive quality audit, not an EDA-based predictor decision using holdout outcomes. Missing counts are 669, 728, 489, 82 and 99 in 2010–2014 respectively. Pollution was recorded at the US Embassy and weather at Beijing Capital International Airport, introducing a spatial mismatch. Dataset: https://archive.ics.uci.edu/dataset/381/beijing+pm2+5+data ; DOI 10.24432/C5JS49, CC BY 4.0. Environmental context: Liang et al. (2015), DOI 10.1098/rspa.2015.0257.")

    # 3. Additional EDA slide replaces unsupported intuition with measured structure.
    d.start("02 / What training-only EDA tells us", "Recent history is the strongest starting point", "All three graphs use 2010–2012 only; associations guide candidate features, not causal claims.", source="SOURCE: Training autocorrelation, hourly profiles and Pearson correlations; saved EDA outputs.")
    panels = [(56,"Temporal dependence","training_autocorrelation"), (454,"Average daily pattern","training_hourly_profile"), (852,"Weather associations","training_weather_correlations")]
    for x, title, graphic in panels:
        d.panel(x, 209, 372, 366, title)
        d.figure(graphic, x+12, 260, 348, 300)
    for x, title, detail in [(56,"Take persistence seriously","1-hour ACF 0.961; 24-hour 0.394.\nA weekly dependency is not established."), (454,"Encode the clock cyclically","The pooled mean peaks at 01:00.\nAverages do not predict each event."), (852,"Validate the extra information","Wind r = −0.264; dew point +0.245.\nCorrelation alone does not imply cause.")]:
        d.rect(x, 591, 372, 64, SOFT, 8)
        d.text(title, x+13, 599, 346, 17, GREEN, True)
        d.text(detail, x+13, 624, 346, 13, MUTED)
    d.note("EDA interpretation and why the features are plausible", "The lag autocorrelations are measured on the regular hourly training grid without deleting gaps. One-hour autocorrelation is 0.961, 24-hour is 0.394 and 168-hour is approximately −0.006. Strong short-term dependence motivates persistence and recent lags. It does not prove each proposed lag is useful; grouped ablation later tests additions. The hourly graph pools the 2010–2012 values by hour of day. The highest mean is at 01:00 (113.08, based on 1,020 observations). October has the largest monthly mean (118.17) in the saved training profile, though the deck focuses on the daily graph for readability. Calendar features are candidates, not a claim that every day repeats yesterday. Numeric weather correlations with PM2.5 include wind speed −0.264, dew point +0.245, pressure −0.138, rain hours −0.055, snow hours +0.011 and temperature +0.003. Small Pearson correlations do not exclude nonlinear predictive effects. Spatial mismatch and correlated weather variables prevent a causal interpretation. The right tail and 7.17% training missingness motivate retaining peaks and reporting historical support. These findings justify features and diagnostics before inspecting the holdout, while chronological validation determines whether additions improve forecasts.")

    # 4. Forecast timing, complete feature groups and theoretical model comparison.
    d.start("03 / Features and model formulation", "Build one forecast row without future information", "Predict y(t + 1) from measured history through t · 70 raw columns → 73 with the four observed wind categories.")
    d.rect(56, 208, 1168, 72, NAVY, 10)
    d.text("ORIGIN  t = 10:00", 77, 222, 310, 20, MINT, True)
    d.text("Observed PM2.5 + weather + trailing history", 77, 251, 630, 17, WHITE)
    d.text("→", 757, 224, 65, 36, MINT, True)
    d.text("TARGET  t + 1 = 11:00", 860, 222, 343, 20, WHITE, True)
    d.text("Only the future calendar is already known", 860, 251, 343, 15, "BED1D2")
    d.panel(56, 301, 572, 303, "Four feature groups")
    features=[("20", "Pollution lags + missingness", "Offsets 0, 1, 2, 3, 6, 12, 24, 48, 72, 168 h"), ("18", "Trailing statistics", "3, 6, 12, 24, 72, 168 h: mean, std, observed fraction"), ("7", "Known target calendar", "Hour / weekday / month sine–cosine + weekend"), ("25", "Weather + wind direction", "6 weather fields at 0 / 1 / 24 h; flags + wind encoding")]
    for i,(count,title,detail) in enumerate(features):
        yy=350+i*60
        d.text(count,75,yy,48,24,GREEN,True)
        d.text(title,140,yy,464,18,INK,True)
        d.text(detail,140,yy+27,464,13,MUTED)
        if i<3: d.line(75,yy+52,607,yy+52,LINE,0.7)
    d.panel(652,301,572,303,"Four models, one target")
    models=[("Persistence","ŷ(t + 1) = y(t)","Strong latest-reading reference"), ("Daily seasonal naive","ŷ(t + 1) = y(t − 23)","Same target hour on the previous day"), ("Ridge regression","Squared error + L2 penalty (α = 10)","Fold-fitted median imputation and scaling"), ("XGBoost","Additive trees + complexity penalty","Native numeric missing splits; outputs clipped ≥ 0")]
    for i,(name,formula,detail) in enumerate(models):
        yy=350+i*60
        d.text(name,672,yy,220,18,INK,True)
        d.text(formula,904,yy,300,15,GREEN)
        d.text(detail,672,yy+27,532,14,MUTED)
        if i<3: d.line(672,yy+52,1204,yy+52,LINE,0.7)
    d.text("Safeguards: 168-hour warm-up · windows end at t · fold-fitted preprocessing · evaluation targets stay observed",60,629,1150,17,GREEN,True)
    d.note("Feature accounting, lag zero and model theory", "The ten PM2.5 offsets are 0, 1, 2, 3, 6, 12, 24, 48, 72, 168, each with a missingness flag (20 columns). Lag zero is the current observed y(t), not the future label. Six trailing windows each supply mean, population standard deviation (ddof 0) and observed fraction (18). Missing summaries remain missing when the support is empty. Seven calendar columns use sine/cosine for known target hour, weekday and month plus weekend. Six numeric weather fields at 0, 1 and 24 hours contribute 18 columns, current weather missing flags contribute six, and current wind contributes one categorical column. Total 70; four observed wind categories give 73 encoded columns. Encoders and Ridge medians/scalers are learned inside each fold. Persistence repeats y(t); the seasonal source is t−23, which is one day before the target t+1. Using t−24 would be the wrong clock hour. Ridge fits squared error plus an L2 penalty with alpha 10. XGBoost builds F_K(x)=b+sum eta f_k(x), penalizing tree leaves and leaf weights; gradients/Hessians yield leaf score −G/(H+lambda). Depth, child weight, shrinkage, sampling and regularization control complexity. Learned outputs are clipped at zero consistently in all phases. Sources: Chen and Guestrin (2016), https://arxiv.org/abs/1603.02754 ; Hoerl and Kennard (1970), DOI 10.1080/00401706.1970.10488634.")

    # 5. Quarterly time order and a real validation feature-comparison plot.
    d.start("04 / Chronological design and selection", "Tune on 2013, then freeze the settings", "Quarterly expanding refits · observed histories update hourly within each evaluation block.",source="SOURCE: configs/phase1.toml, selection.json and validation_ablation.csv; rolling-origin forecasting principles.")
    for x,w,year,label,color in [(56,372,"2010–2012","Initial training + EDA",GREEN),(454,372,"2013","Q1   Q2   Q3   Q4 · validation",BLUE),(852,372,"2014","Q1   Q2   Q3   Q4 · holdout",NAVY)]:
        d.rect(x,207,w,97,color,10)
        d.text(year,x+19,220,w-38,30,WHITE,True)
        d.text(label,x+19,264,w-38,17,WHITE)
    d.panel(56,328,572,324,"Does extra information help XGBoost?","Selected parameters fixed; validation dates only")
    d.figure("validation_ablation",70,408,544,218)
    d.text("Features improve XGBoost; simpler validation models remain stronger.",76,629,531,13,ORANGE,True)
    d.panel(652,328,572,324,"Selection and label-availability rules")
    d.text("Three-candidate grid → lowest pooled validation RMSE",672,382,531,18,INK,True)
    d.text("Candidate RMSE: 26.772 / 26.801 / 26.953",672,416,531,17,MUTED)
    d.text(f"Selected: {selected['n_estimators']} trees · depth {selected['max_depth']} · learning rate {selected['learning_rate']}",672,454,531,19,GREEN,True)
    d.text("Child weight 5 · row 0.8 / column 0.9 · L2 = 5",672,485,531,17,MUTED)
    d.line(672,520,1204,520,LINE)
    d.text("For target block B: latest training label ≤ B − 1 h",672,538,531,18,INK,True)
    d.text("No holdout tuning or early stopping; fit preprocessing per fold.\nSame observed targets and naive inputs for all models.",672,579,531,16,MUTED,leading=1.5)
    d.note("Time splits, parameter choice, ablation and fair coverage", "The 2010–2012 period is initial training and EDA. Four expanding quarterly folds in 2013 choose parameters using pooled validation RMSE. Four quarterly 2014 folds then evaluate the frozen setting. The figure shows XGBoost validation RMSE falling 27.513 → 27.092 → 26.772 as trailing/calendar and weather groups are added. The selected settings remain fixed during this comparison; it is not separate optimal tuning for every group. Dashed reference lines show persistence 24.870 and Ridge 23.809, both better than full XGBoost on validation. Candidate RMSEs are 26.772, 26.801, 26.953, a narrow separation rather than a globally proven optimum. Selected XGBoost has 300 trees, depth 3, learning rate 0.05, child weight 5, subsample 0.8, colsample 0.9 and lambda 5, with seed 42, histogram trees and two threads. There is no test-based early stopping. For block start B, the first forecast origin is B−h and every training label must be known by that origin. Q1 2014 uses labels through 31 December 2013 at 23:00, with 32,951 eligible training rows; Q4 has 39,447. Later observed quarters can enter later fits under the same policy. The shared mask requires warm-up, a real target and both naive inputs, yielding 8,563 validation and 8,531 holdout targets. Exclusion counts overlap. Methodological basis: https://otexts.com/fpp3/toolbox.html")

    # 6. A labelled bar chart, full metric table and explicit skill/coverage values.
    d.start("05 / Pooled holdout comparison", "The gain is useful, but modest", "Same 8,531 observed 2014 target hours for all four models · MAE / RMSE in µg/m³; R² is dimensionless.", source="SOURCE: Real-data metrics.csv; pooled errors recomputed from saved out-of-fold predictions.")
    d.stat(56,199,372,f"{improvement['xgboost']:.2f}%","XGBoost RMSE improvement","Relative to persistence, on common hours")
    d.stat(454,199,372,f"{improvement['ridge']:.2f}%","Ridge RMSE improvement","The linear comparator is slightly better")
    d.stat(852,199,372,f"{100*n/8760:.2f}%","Holdout target coverage","8,531 / 8,760 expected hours; 229 excluded")
    d.panel(56,310,620,310,"Annual RMSE comparison")
    d.figure("holdout_rmse",68,354,596,253)
    d.panel(700,310,524,310,"All required metrics")
    for x,label,w in [(721,"Model",208),(935,"MAE",80),(1035,"RMSE",84),(1140,"R²",62)]:
        d.text(label,x,365,w,16,MUTED,True)
    for i,(key,label) in enumerate([("persistence","Persistence"),("seasonal_naive","Seasonal naive"),("ridge","Ridge"),("xgboost","XGBoost")]):
        yy=405+i*46
        if key=="ridge": d.rect(713,yy-7,498,37,SOFT,6)
        color=GREEN if key=="ridge" else INK
        d.text(label,721,yy,205,17,color,key=="ridge")
        for x,col,w in [(935,"mae",82),(1035,"rmse",91),(1140,"r2",74)]:
            places=4 if col=="r2" else 3
            d.text(f"{float(scores[key][col]):.{places}f}",x,yy,w,17,color,key=="ridge")
    d.text("Ridge has lower RMSE by 0.025 µg/m³. The significance of this difference has not been tested.",60,637,1154,18,GREEN,True)
    d.note("Pooled results, metrics and interpretation", "All four methods are scored on the same 8,531 observed target hours out of 8,760 in 2014: 97.39% coverage. The 229 excluded hours lack a target or at least one naive input; the exclusion counts overlap. The graph starts at zero and shows the weak seasonal reference without hiding its large error. Persistence MAE 11.926, RMSE 22.033 and R² 0.9442; daily seasonal naive 67.296, 98.754 and −0.1205; Ridge 11.825, 21.063 and 0.9490; XGBoost 11.833, 21.088 and 0.9489. XGBoost improves RMSE 4.29% over persistence, while Ridge improves 4.41%. The learned-model gap of about 0.025 is not a tested statistically significant difference. MAE averages absolute errors; RMSE square-roots the mean squared error and emphasizes large misses. R² compares squared errors with the observed-period mean denominator, may be negative and must not be called percentage accuracy. The negative seasonal R² is retained. Pooled metrics use all errors rather than averaged quarterly RMSEs. XGBoost's favorable 2014 comparison does not erase its weaker 2013 validation comparison. Historical point errors do not give calibrated intervals for future predictions.")

    # 7. Dedicated full-size seasonality and high-pollution comparisons.
    d.start("06 / Stability and high-pollution diagnostics", "A single annual score hides important differences", "Quarterly model rankings change, and XGBoost's pooled advantage does not extend to the high-pollution subset.", source="SOURCE: fold_metrics.csv and errors_by_pollution_regime.csv; high threshold fixed at training p90 = 218 µg/m³.")
    d.panel(56,210,572,326,"Quarterly holdout RMSE")
    d.figure("quarterly_rmse",69,257,546,266)
    d.panel(652,210,572,326,"Errors by observed pollution level")
    d.figure("pollution_regime_errors",665,257,546,266)
    d.rect(56,553,572,99,SOFT,9)
    d.text("Both learned models beat persistence each quarter.",76,567,532,19,GREEN,True)
    d.text("Ridge leads Q1 / Q3; XGBoost leads Q2 / Q4.\nIn Q1, XGBoost's MAE still exceeds persistence.",76,602,532,17,MUTED,leading=1.35)
    d.rect(652,553,572,99,SOFT,9)
    d.text("High-pollution RMSE: XGBoost 39.959 vs 38.557.",672,567,532,19,ORANGE,True)
    d.text("High-pollution bias: XGBoost −11.912; Ridge −13.313.\nBoth learned models underpredict this subset on average.",672,602,532,16,MUTED,leading=1.4)
    d.note("Quarterly variation and pollution regimes", "Quarterly RMSE makes temporal variation visible. Persistence is 27.298, 22.208, 14.527 and 22.255; Ridge 26.063, 21.311, 14.012 and 21.148; XGBoost 26.839, 20.758, 14.087 and 20.770. Both learned methods beat persistence on RMSE each quarter, but Ridge leads Q1 and Q3 and XGBoost Q2 and Q4. XGBoost's Q1 MAE is worse than persistence despite the smaller RMSE. We fix the high-pollution threshold at 218 using the 2010–2012 p90, without consulting the holdout distribution. There are 865 high and 7,666 lower-pollution common hours. On high hours, persistence RMSE is 38.557, Ridge 38.522 and XGBoost 39.959. Below 218, errors are 19.301, 18.062 and 17.740, so the pooled XGBoost gain is largely associated with the lower-pollution subset. High-pollution bias is −5.534 for persistence, −13.313 for Ridge and −11.912 for XGBoost. Negative bias means underprediction. These are descriptive subsets, not a new optimized experiment. No significant ranking or causal mechanism is inferred from the differences. A practical peak-focused use case needs a separately evaluated peak-error objective.")

    # 8. Actual forecasting behavior, all-hour residuals and the largest failure.
    d.start("07 / Forecast trace and failure analysis", "Look at the predictions, not just the summary", "Saved rolling backtest forecasts · the final deployment refit is never used to generate these historical scores.",source="SOURCE: 2014 predictions.csv; trace shows 25–31 December; the spike is the largest XGBoost holdout miss.")
    d.panel(56,199,1168,241,"Actual versus predicted · final seven days of 2014")
    d.figure("forecast_last_week",70,243,1140,188)
    d.panel(56,458,572,195,"Signed residuals across the full holdout")
    d.figure("residual_distribution",66,498,552,151)
    d.panel(652,458,572,195,"Largest spike miss · 9 April 2014, 20:00")
    d.figure("largest_spike_miss",662,498,552,151)
    residual_scope=chart_audit["residual_plot"]
    d.note("Forecast trace, error distribution and abrupt spike", f"The upper graph plots every common forecast hour in the final seven days of 2014, with missing scored hours represented as breaks. Observed concentration, Ridge and XGBoost share the same timestamps. These are saved out-of-fold predictions made by quarterly fitted models, not the final archive-trained models. Tracking the recent level does not guarantee anticipating a sharp change. The lower-left graph uses signed residuals prediction minus observation from all 8,531 holdout hours. Its display is restricted to ±80 µg/m³ for legibility: {residual_scope['xgboost']['residual_view_outside_80']} XGBoost and {residual_scope['ridge']['residual_view_outside_80']} Ridge errors lie outside this view; histogram percentages retain the denominator of all 8,531 hours. Thus the tails are not silently removed from metrics. Negative residuals are underpredictions. XGBoost median absolute error is 6.88, p90 26.28 and p95 37.46; these empirical quantiles are not prediction intervals. One-hour residual autocorrelation is 0.120 for XGBoost and 0.063 for Ridge, so residual structure remains. The lower-right graph is the largest XGBoost miss, chosen explicitly as a failure case. PM2.5 rises from 80 to 580 in one hour; persistence predicts 80, Ridge 92.469, XGBoost 86.801, an XGBoost absolute miss of about 493.199. Without an independent sensor-quality flag, we retain this target. The visual example explains a failure mode rather than proving all peaks have one cause.")

    # 9. Aircast and reproducibility, with final-refit importance clearly scoped.
    d.start("08 / Implementation, explanation and next steps", "From experiment artifacts to the Aircast demo", "The interface invokes the original Python feature builder and saved pipelines; training and inference are separate.",source="SOURCE: Saved deployment-refit gain importance and the recorded historical/local prediction audit.")
    for i,(x,label,body) in enumerate([(56,"INPUT","Manual observation / hourly CSV"),(360,"VALIDATE","Flask + revalidated history"),(664,"PREDICT","Saved Ridge / XGBoost pipeline"),(968,"EXPLORE","History chart + forecast export")]):
        d.rect(x,207,256,87,NAVY,9)
        d.text(label,x+16,219,224,13,MINT,True)
        d.text(body,x+16,248,224,17,WHITE)
        if i<3: d.text("→",x+266,233,30,25,GREEN,True)
    d.panel(56,319,572,333,"What the final XGBoost refit uses","Top five features · gain importance is not causation")
    d.figure("deployment_feature_importance",68,397,548,233)
    d.panel(652,319,572,333,"Reproducibility and practical limits")
    d.text("17,062",672,370,207,39,GREEN,True)
    d.text("refitted predictions reproduced",891,379,310,18,INK,True)
    d.text("Maximum numeric difference < 6e−14 in the recorded audit.",672,422,532,16,MUTED)
    d.line(672,453,1204,453,LINE)
    d.text("Retained: config, hashes, versions, fold audits, predictions,\ncleaning exports, feature dictionary, CLI / notebook and tests.",672,472,532,16,INK,leading=1.45)
    d.text("Limits: historical Beijing; peak errors; assumed data availability.\nNext: fresh evaluation, peak-focused methods and uncertainty.",672,544,532,16,MUTED,leading=1.45)
    d.text("Demo scenarios are historical. No calibrated intervals or AQI mapping.",672,616,532,13,ORANGE,True)
    d.note("Aircast, gain importance, audits and future work", "Aircast lets a user append a manual reading or upload an hourly CSV, select an available model, inspect history, generate a next-hour forecast, export results and open the report. The frontend is HTML/CSS/JavaScript with an SVG chart, and Flask validates data before invoking the original feature builder. Uploads need at least 169 hourly slots, unique valid timestamps and physical numeric values, and are bounded at 4 MB / 100,000 rows or hourly slots. The normalized recent window travels with each request and is revalidated, so different server instances can forecast without shared upload sessions. Final model artifacts are refitted only after evaluation and reject histories ending before their fit endpoint. Gain importance belongs to that final refit: latest PM2.5 51.1% and three-hour mean 24.5%; it is not causal attribution, a stability study or holdout feature selection. The recorded independent audit refitted both learned methods across four holdout quarters, reproducing all 17,062 predictions with maximum floating-point difference below 6e−14. It also checked actual targets, references, recomputed metrics, manual-input response and stateless upload parity. Data/source hashes, configuration, versions, selection, fold audits, predictions, model artifacts, preparation exports and tests remain inspectable. Synthetic fixtures test behavior and do not supply scores. Vercel reported Ready; the existing record does not independently verify hosted inference or browser rendering. Sample data is historical. The original uploaded archive gives 2015-01-01 00:00 predictions XGBoost 11.92273 and Ridge 13.63978, unscored because the target is outside the archive. No AQI mapping or calibrated prediction interval is supplied. Future peak objectives, uncertainty methods and sequence/hybrid models require fresh evaluation or an exploratory label after inspecting 2014.")

    # 10. Requested closing slide, retaining the exact team details and no emails.
    d.start("Team / Questions and discussion",dark=True)
    d.text("Thank you.",56,115,1120,79,WHITE,True)
    d.text("Questions & discussion",60,220,1100,27,MINT)
    d.text("TEAM MEMBER",62,306,670,15,"BED1D2",True)
    d.text("ENROLLMENT NUMBER",823,306,360,15,"BED1D2",True)
    for i,(name,enrollment) in enumerate(TEAM):
        yy=350+i*51
        d.text(name,62,yy,670,27,WHITE,True)
        d.text(enrollment,823,yy+1,360,25,MINT)
        if i<4: d.line(62,yy+39,1199,yy+39,"2C4753",0.8)
    d.text("Phase 1: a reproducible pipeline, modest pooled gains, and clear evidence for the next experiment.",62,628,1140,18,"BED1D2")
    d.note("Thank you, team and viva prompts", "Thank the audience and invite questions. Team: Samarth Khera — 2401010407; Pratham Malhotra — 2401010343; Adityaraj Pal — 2401010042; Tattva Rajput — 2401010478; Himanshu — 2401010188. No email addresses or invented individual contribution allocations are included. Phase 1 delivers a complete classical ML workflow and Aircast demonstration. The measured conclusion is modest 2014 pooled improvement beyond persistence, a close Ridge/XGBoost comparison and unresolved peak performance. Explain lag zero, t−23 seasonal alignment, training-label availability, common scoring hours, the weak 2013 validation comparison and the difference between final refits and historical backtests. The first common target is 2014-01-01 00:00: origin latest PM2.5 23, previous reading 20, three-hour mean 22, observed target 24, persistence 23, Ridge 26.688, XGBoost 21.078 and seasonal 84. It is chosen by time rather than model success and shows why an individual forecast does not determine a year-long ranking. New model choices motivated by inspected holdout errors need fresh evaluation. References and derivations are retained in the seven-page report, the slide notes and the Phase 1 literature review.")
    d.save()
    print(f"Built {d.page} slides with {len(chart_audit['figures'])} data-driven graph panels")
    print(OUTPUT/"Aircast_Phase1_Presentation.pptx")
    print(OUTPUT/"slides.pdf")


if __name__=="__main__":
    main()
